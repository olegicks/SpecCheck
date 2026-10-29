import json

import httpx
import pytest
from openpyxl import load_workbook

from app.agent.fake_llm import ScriptedLLM, respond, tool_call
from tests.api.conftest import make_settings, run_check
from tests.conftest import make_pdf

PNG_MAGIC = b"\x89PNG"


def parse_sse(text: str) -> list[tuple[str, dict]]:
    events = []
    for block in text.split("\n\n"):
        fields = dict(
            line.split(": ", 1)
            for line in block.splitlines()
            if ": " in line and not line.startswith(":")
        )
        if "event" in fields:
            events.append((fields["event"], json.loads(fields["data"])))
    return events


async def test_health_reports_offline_mode(client) -> None:
    body = (await client.get("/health")).json()
    assert body["status"] == "ok"
    assert body["mode"] == "offline"
    assert body["model"] == "offline-demo"


async def test_health_reports_live_mode(app_factory) -> None:
    app = app_factory(make_settings(OPENAI_API_KEY="sk-test-" + "a" * 30))
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t") as c:
        body = (await c.get("/health")).json()
    assert body["mode"] == "live" and body["model"] == "gpt-4o-mini"


async def test_samples_and_templates(client) -> None:
    samples = (await client.get("/api/samples")).json()
    assert samples[0]["id"] == "sc-mcu-100" and samples[0]["page_count"] == 8
    templates = (await client.get("/api/templates")).json()
    assert templates[0]["id"] == "vehicle-computer-mcu"
    assert len(templates[0]["requirements"]) == 9


# --- requirements import/export ----------------------------------------------------


async def test_requirements_import_export_round_trip(client, template_requirements) -> None:
    for fmt in ("yaml", "csv"):
        exported = await client.post(
            "/api/requirements/export", json={"format": fmt, "requirements": template_requirements}
        )
        assert exported.status_code == 200
        imported = await client.post(
            "/api/requirements/import", json={"format": fmt, "content": exported.text}
        )
        assert imported.status_code == 200, imported.text
        assert imported.json() == template_requirements


async def test_requirements_import_errors(client) -> None:
    response = await client.post(
        "/api/requirements/import", json={"format": "csv", "content": "id,title\nREQ-1,x"}
    )
    assert response.status_code == 422
    assert "missing column" in response.json()["detail"]["errors"][0]


# --- datasheets ---------------------------------------------------------------------


async def test_upload_valid_pdf(client) -> None:
    pdf = make_pdf(["<p>Operating temperature -40 to 125 °C</p>"])
    response = await client.post(
        "/api/datasheets", files={"file": ("my.pdf", pdf, "application/pdf")}
    )
    assert response.status_code == 201
    body = response.json()
    assert body["filename"] == "my.pdf" and body["page_count"] == 1 and not body["is_sample"]
    assert (await client.get(f"/api/datasheets/{body['id']}")).status_code == 200


async def test_upload_wrong_file_type(client) -> None:
    response = await client.post(
        "/api/datasheets", files={"file": ("evil.pdf", b"MZ\x90\x00 executable", "application/pdf")}
    )
    assert response.status_code == 422
    assert "not a PDF" in response.json()["detail"]


async def test_upload_too_large(app_factory) -> None:
    app = app_factory(make_settings(SPECCHECK_MAX_UPLOAD_MB=1))
    data = b"%PDF-1.7\n" + b"0" * (1024 * 1024 + 10)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t") as c:
        response = await c.post("/api/datasheets", files={"file": ("big.pdf", data)})
    assert response.status_code == 413


async def test_upload_too_many_pages(app_factory) -> None:
    app = app_factory(make_settings(SPECCHECK_MAX_PDF_PAGES=2))
    pdf = make_pdf(["<p>a</p>", "<p>b</p>", "<p>c</p>"])
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t") as c:
        response = await c.post("/api/datasheets", files={"file": ("x.pdf", pdf)})
    assert response.status_code == 422 and "limit is 2" in response.json()["detail"]


async def test_upload_rate_limit(app_factory) -> None:
    app = app_factory(make_settings(SPECCHECK_RATE_LIMIT_UPLOADS="2/minute"))
    pdf = make_pdf(["<p>hello world</p>"])
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t") as c:
        codes = [
            (await c.post("/api/datasheets", files={"file": ("x.pdf", pdf)})).status_code
            for _ in range(3)
        ]
    assert codes == [201, 201, 429]


async def test_unknown_datasheet(client) -> None:
    assert (await client.get("/api/datasheets/ds_nope")).status_code == 404
    assert (await client.get("/api/datasheets/ds_nope/pages/1.png")).status_code == 404


async def test_page_png_with_highlight(client) -> None:
    response = await client.get(
        "/api/datasheets/sc-mcu-100/pages/5.png", params={"highlight": "I/O supply voltage"}
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "image/png"
    assert response.content.startswith(PNG_MAGIC)
    assert (await client.get("/api/datasheets/sc-mcu-100/pages/99.png")).status_code == 404


# --- checks ----------------------------------------------------------------------------


async def test_check_validation_errors(client, template_requirements) -> None:
    r = await client.post("/api/checks", json={"datasheet_id": "sc-mcu-100", "requirements": []})
    assert r.status_code == 422
    dup = [template_requirements[0], template_requirements[0]]
    r = await client.post("/api/checks", json={"datasheet_id": "sc-mcu-100", "requirements": dup})
    assert r.status_code == 422
    r = await client.post(
        "/api/checks", json={"datasheet_id": "ds_x", "requirements": template_requirements}
    )
    assert r.status_code == 404
    bad = [{**template_requirements[0], "type": "range", "parameters": {"min": 1}}]
    r = await client.post("/api/checks", json={"datasheet_id": "sc-mcu-100", "requirements": bad})
    assert r.status_code == 422
    assert any("max" in e["msg"] for e in r.json()["detail"])


async def test_check_rate_limit(app_factory, template_requirements) -> None:
    app = app_factory(make_settings(SPECCHECK_RATE_LIMIT_CHECKS="2/minute"))
    body = {"datasheet_id": "sc-mcu-100", "requirements": template_requirements[:1]}
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t") as c:
        codes = [(await c.post("/api/checks", json=body)).status_code for _ in range(3)]
        await app.state.service.wait_all()
    assert codes == [202, 202, 429]


async def test_concurrent_run_cap(app_factory, template_requirements) -> None:
    app = app_factory(
        make_settings(SPECCHECK_MAX_CONCURRENT_RUNS=1),
        llm_factory=lambda reqs, ds: ScriptedLLM([], delay=0.2),
    )
    body = {"datasheet_id": "sc-mcu-100", "requirements": template_requirements[:1]}
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t") as c:
        first = await c.post("/api/checks", json=body)
        second = await c.post("/api/checks", json=body)
        await app.state.service.wait_all()
    assert first.status_code == 202 and second.status_code == 429


async def test_check_lifecycle_sse_override_and_exports(client, template_requirements) -> None:
    check_id = await run_check(client, template_requirements)

    # result
    run = (await client.get(f"/api/checks/{check_id}")).json()
    assert run["state"] == "completed"
    assert run["mode"] == "offline" and run["model"] == "offline-demo"
    assert run["summary"] == {"total": 9, "passed": 7, "failed": 1, "unclear": 1, "overridden": 0}

    # SSE replay
    sse = await client.get(f"/api/checks/{check_id}/events")
    assert sse.headers["content-type"].startswith("text/event-stream")
    events = parse_sse(sse.text)
    kinds = [e[1]["type"] for e in events if e[0] == "agent"]
    assert kinds[0] == "run_started" and kinds[-1] == "run_completed"
    assert {"search", "read_page", "submit", "verify", "decide"} <= set(kinds)
    assert events[-1] == ("done", {"state": "completed", "id": check_id})
    seqs = [e[1]["seq"] for e in events if e[0] == "agent"]
    assert seqs == list(range(1, len(seqs) + 1))

    # resume with Last-Event-ID
    resumed = await client.get(
        f"/api/checks/{check_id}/events", headers={"Last-Event-ID": str(seqs[-3])}
    )
    assert len([e for e in parse_sse(resumed.text) if e[0] == "agent"]) == 2

    # override needs a comment
    url = f"/api/checks/{check_id}/findings/REQ-005"
    assert (await client.patch(url, json={"status": "PASS", "comment": ""})).status_code == 422
    assert (
        await client.patch(url, json={"status": "MAYBE", "comment": "ok ok"})
    ).status_code == 422
    r = await client.patch(url, json={"status": "PASS", "comment": "We order the 100B variant."})
    assert r.status_code == 200
    body = r.json()
    assert body["status"] == "PASS" and body["agent_status"] == "UNCLEAR"  # original kept
    assert body["override"]["comment"] == "We order the 100B variant."
    assert not body["needs_review"]
    run = (await client.get(f"/api/checks/{check_id}")).json()
    assert run["summary"]["passed"] == 8 and run["summary"]["overridden"] == 1
    unknown = await client.patch(
        f"/api/checks/{check_id}/findings/REQ-999", json={"status": "PASS", "comment": "abc"}
    )
    assert unknown.status_code == 404

    # exports
    md = await client.get(f"/api/checks/{check_id}/export", params={"format": "md"})
    assert md.status_code == 200 and "# SpecCheck report" in md.text
    assert "Reviewer override" in md.text and "attachment" in md.headers["content-disposition"]
    csv_text = (await client.get(f"/api/checks/{check_id}/export?format=csv")).text
    assert csv_text.lstrip("﻿").startswith("Requirement,Title")
    assert csv_text.count("\n") == 10
    data = (await client.get(f"/api/checks/{check_id}/export?format=json")).json()
    assert data["id"] == check_id and len(data["results"]) == 9
    xlsx = await client.get(f"/api/checks/{check_id}/export?format=xlsx")
    import io

    wb = load_workbook(io.BytesIO(xlsx.content))
    assert wb["Results"].max_row == 10 and wb["Run"]["B1"].value == check_id
    assert (await client.get(f"/api/checks/{check_id}/export?format=pdf")).status_code == 422

    # clear override
    cleared = await client.delete(f"/api/checks/{check_id}/findings/REQ-005/override")
    assert cleared.json()["status"] == "UNCLEAR"


async def test_unknown_check(client) -> None:
    assert (await client.get("/api/checks/chk_nope")).status_code == 404
    assert (await client.get("/api/checks/chk_nope/events")).status_code == 404
    assert (await client.get("/api/checks/chk_nope/export")).status_code == 404
    r = await client.patch(
        "/api/checks/chk_nope/findings/REQ-1", json={"status": "PASS", "comment": "abc"}
    )
    assert r.status_code == 404
    assert (await client.delete("/api/checks/chk_nope/findings/REQ-1/override")).status_code == 404


async def test_running_check_rejects_override_and_export(
    app_factory, template_requirements
) -> None:
    app = app_factory(llm_factory=lambda reqs, ds: ScriptedLLM([], delay=0.3))
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t") as c:
        r = await c.post(
            "/api/checks",
            json={"datasheet_id": "sc-mcu-100", "requirements": template_requirements[:1]},
        )
        check_id = r.json()["id"]
        assert r.json()["state"] == "running"
        override = await c.patch(
            f"/api/checks/{check_id}/findings/REQ-001", json={"status": "PASS", "comment": "abc"}
        )
        export = await c.get(f"/api/checks/{check_id}/export")
        await app.state.service.wait_all()
    assert override.status_code == 409 and export.status_code == 409


async def test_check_on_uploaded_pdf_offline_is_unclear(client, template_requirements) -> None:
    pdf = make_pdf(["<p>Operating temperature -40 to 125 °C</p>"])
    ds = (await client.post("/api/datasheets", files={"file": ("u.pdf", pdf)})).json()
    check_id = await run_check(client, template_requirements[:2], ds["id"])
    run = (await client.get(f"/api/checks/{check_id}")).json()
    assert run["summary"]["unclear"] == 2
    assert "OPENAI_API_KEY" in run["results"][0]["reason"]


async def test_agent_crash_marks_run_failed(app_factory, template_requirements) -> None:
    def broken(_messages):
        raise RuntimeError("unexpected")

    app = app_factory(llm_factory=lambda reqs, ds: ScriptedLLM([broken]))
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t") as c:
        check_id = await run_check(c, template_requirements[:2])
        run = (await c.get(f"/api/checks/{check_id}")).json()
        events = parse_sse((await c.get(f"/api/checks/{check_id}/events")).text)
    assert run["state"] == "failed"
    assert run["error"] == "Internal error: RuntimeError"
    assert all(r["status"] == "UNCLEAR" for r in run["results"])
    assert events[-1] == ("done", {"state": "failed", "id": check_id})


@pytest.mark.parametrize("path", ["/health", "/api/samples"])
async def test_cors_allows_frontend_origin(client, path: str) -> None:
    r = await client.get(path, headers={"Origin": "http://localhost:5173"})
    assert r.headers["access-control-allow-origin"] == "http://localhost:5173"


async def test_unhandled_errors_are_generic(app_factory) -> None:
    app = app_factory()

    @app.get("/boom")
    async def boom():
        raise RuntimeError("secret internals")

    transport = httpx.ASGITransport(app=app, raise_app_exceptions=False)
    async with httpx.AsyncClient(transport=transport, base_url="http://t") as c:
        r = await c.get("/boom")
    assert r.status_code == 500 and r.json() == {"detail": "Internal server error."}


async def test_scripted_live_like_run(app_factory, template_requirements) -> None:
    """A run driven by a scripted 'model' (stand-in for OpenAI) produces the same pipeline."""
    finding = {
        "requirement_id": "REQ-006", "value": 3, "unit": None, "page": 3, "confidence": 0.95,
        "quote": "The device contains 3 CAN FD controllers (MCAN0 to MCAN2).", "notes": None,
    }  # fmt: skip
    app = app_factory(
        llm_factory=lambda reqs, ds: ScriptedLLM(
            [respond(tool_call("submit_finding", finding)), respond(content="Done.")],
            model="gpt-4o-mini",
        )
    )
    req = next(r for r in template_requirements if r["id"] == "REQ-006")
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t") as c:
        check_id = await run_check(c, [req])
        run = (await c.get(f"/api/checks/{check_id}")).json()
    assert run["model"] == "gpt-4o-mini"
    assert run["results"][0]["status"] == "FAIL"
    assert run["results"][0]["reason"] == "3 × CAN FD controller < required 4"
